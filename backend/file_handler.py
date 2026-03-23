import os
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from PIL import Image
from backend.config import settings

# Safe imports
try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    import google.generativeai as genai
    if settings.gemini_api_key:
        genai.configure(api_key=settings.gemini_api_key)
    GENAI_AVAILABLE = True
except (ImportError, TypeError):
    genai = None
    GENAI_AVAILABLE = False


# Ensure upload directory exists
BASE_UPLOAD_DIR = Path(settings.upload_dir)
BASE_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

def save_upload_to_disk(filename: str, file_bytes: bytes, user_id: str) -> Path:
    """Saves the uploaded file to a user-specific folder."""
    user_folder = BASE_UPLOAD_DIR / str(user_id)
    user_folder.mkdir(parents=True, exist_ok=True)
    
    # Sanitize filename
    sanitized_name = filename.replace("/", "_").replace("\\", "_")
    full_path = user_folder / sanitized_name
    
    with open(full_path, "wb") as f:
        f.write(file_bytes)
        
    return full_path

def extract_text_from_pdf(path: Path) -> list:
    try:
        reader = PdfReader(str(path))
        pages = []
        for i, p in enumerate(reader.pages):
            text = p.extract_text()
            if text and text.strip():
                pages.append({"text": text.strip(), "page": i + 1})
        return pages
    except Exception as e:
        print(f"Error reading PDF {path}: {e}")
        return []

def extract_text_from_docx(path: Path) -> list:
    try:
        doc = DocxDocument(str(path))
        # DOCX doesn't have native page numbers easily accessible via python-docx.
        # We'll treat the whole document as page 1, or try to chunk it later.
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        if text.strip():
            return [{"text": text.strip(), "page": 1}]
        return []
    except Exception as e:
        print(f"Error reading DOCX {path}: {e}")
        return []

def extract_text_from_pptx(path: Path) -> list:
    try:
        prs = Presentation(str(path))
        slides = []
        for i, slide in enumerate(prs.slides):
            text = []
            for shape in slide.shapes:
                if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                    text.append(shape.text_frame.text)
            slide_text = "\n".join(text).strip()
            if slide_text:
                slides.append({"text": slide_text, "page": i + 1})
        return slides
    except Exception as e:
        print(f"Error reading PPTX {path}: {e}")
        return []

def extract_text_from_image(path: Path) -> list:
    try:
        # Try Gemini Vision via REST (Better than Tesseract)
        try:
            from backend.core.llm_interface import llm_client
            with open(path, "rb") as img_file:
                img_bytes = img_file.read()
            
            # Determine mime type roughly
            ext = path.suffix.lower()
            mime = "image/png" if ext == ".png" else "image/jpeg"
            
            response = llm_client.generate_with_image(
                prompt="Transcribe all text from this image exactly.", 
                image_bytes=img_bytes,
                mime_type=mime
            )
            
            if "LLM Request Failed" not in response and "Error" not in response:
                return [{"text": f"[Image Transcription]\n{response}", "page": 1}]
            
        except Exception as e:
            print(f"Gemini Vision REST failed: {e}")
        
        # Fallback to Tesseract
        if pytesseract:
            param = Image.open(path)
            # pytesseract expects path or image object
            text = pytesseract.image_to_string(param)
            return [{"text": text, "page": 1}] if text.strip() else []
        else:
            return [{"text": "[Image OCR Failed: Tesseract not installed]", "page": 1}]
            
    except Exception as e:
        print(f"Error reading Image {path}: {e}")
        return [{"text": f"[Image Content - Processing Failed: {e}]", "page": 1}]

def extract_text_from_audio(path: Path) -> list:
    try:
        if not GENAI_AVAILABLE or not settings.gemini_api_key or not genai:
            return [{"text": "[Audio Transcription Failed: API Key missing or Library incompatible]", "page": 1}]
            
        myfile = genai.upload_file(path)
        model = genai.GenerativeModel("gemini-1.5-flash")
        result = model.generate_content(["Generate a transcript of this audio:", myfile])
        return [{"text": f"[Audio Transcript]\n{result.text}", "page": 1}]
    except Exception as e:
        return [{"text": f"[Audio processing failed: {e}]", "page": 1}]

def extract_text_from_file(file_path: Path) -> list:
    """Main entry point to extract text based on extension. Returns list of dicts with text and page num."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()
    pages = []

    if ext == ".pdf":
        pages = extract_text_from_pdf(path)
    elif ext == ".docx":
        pages = extract_text_from_docx(path)
    elif ext in [".pptx", ".ppt"]:
        pages = extract_text_from_pptx(path)
    elif ext == ".txt":
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
            if text.strip():
                pages = [{"text": text.strip(), "page": 1}]
    elif ext in [".jpg", ".jpeg", ".png"]:
        pages = extract_text_from_image(path)
    elif ext in [".mp3", ".wav", ".m4a"]:
        pages = extract_text_from_audio(path)
    else:
        # Try generic read for unknown text types
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
                if text.strip():
                    pages = [{"text": text.strip(), "page": 1}]
        except:
            raise ValueError(f"Unsupported file type: {ext}")

    if not pages:
        return [{"text": f"[File: {path.name} processed but no text found]", "page": 1}]
        
    return pages